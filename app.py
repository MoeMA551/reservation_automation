import os
import re
from datetime import datetime

import pandas as pd
from flask import Flask, render_template, request, send_file
from openpyxl.utils import get_column_letter

UPLOAD_FOLDER = "uploads"
OUTPUT_FOLDER = "outputs"
ALLOWED_EXTENSIONS = {".xlsx", ".csv"}

GROUP_COLUMN = "reservation site"
SUM_COLUMN = "amount"

MAX_SHEET_NAME_LENGTH = 31
INVALID_SHEET_CHARS = r"[:\\/?*\[\]]"
COLUMN_PADDING = 2

app = Flask(__name__)

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)


def read_file(path, extension):
    """Read an uploaded .xlsx or .csv file into a DataFrame."""
    if extension == ".xlsx":
        return pd.read_excel(path)
    try:
        return pd.read_csv(path, encoding="utf-8-sig")
    except UnicodeDecodeError:
        return pd.read_csv(path, encoding="cp1252")


def clean_sheet_name(name):
    """Remove characters Excel doesn't allow and cut to 31 characters."""
    return re.sub(INVALID_SHEET_CHARS, "", str(name))[:MAX_SHEET_NAME_LENGTH]


def autofit_columns(worksheet, data):
    """Set each column's width to fit its longest value or header."""
    for index, column in enumerate(data.columns, start=1):
        longest_value = data[column].fillna("").astype(str).str.len().max()
        width = max(longest_value, len(str(column))) + COLUMN_PADDING
        worksheet.column_dimensions[get_column_letter(index)].width = width


def add_sum_row(worksheet, data, column):
    """Write a SUM formula in the row below the last data row."""
    column_letter = get_column_letter(data.columns.get_loc(column) + 1)
    last_row = len(data) + 1  # +1 for the header row
    worksheet[f"{column_letter}{last_row + 1}"] = (
        f"=SUM({column_letter}2:{column_letter}{last_row})"
    )


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/process", methods=["POST"])
def process_file():
    file = request.files.get("file")
    if file is None or file.filename == "":
        return "Error: no file selected.", 400

    extension = os.path.splitext(file.filename)[1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        return "Error: please upload a .xlsx or .csv file.", 400

    timestamp = datetime.now()
    input_path = os.path.join(UPLOAD_FOLDER, f"{timestamp:%Y%m%d_%H%M%S}{extension}")
    file.save(input_path)

    df = read_file(input_path, extension)

    missing = [col for col in (GROUP_COLUMN, SUM_COLUMN) if col not in df.columns]
    if missing:
        return f"Error: missing column(s): {', '.join(missing)}", 400

    filename = f"{timestamp:%m%d} TL.xlsx"
    output_path = os.path.join(OUTPUT_FOLDER, filename)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for site, site_data in df.groupby(GROUP_COLUMN):
            sheet_name = clean_sheet_name(site)
            site_data.to_excel(writer, sheet_name=sheet_name, index=False)

            worksheet = writer.sheets[sheet_name]
            autofit_columns(worksheet, site_data)
            add_sum_row(worksheet, site_data, SUM_COLUMN)

    return send_file(output_path, as_attachment=True, download_name=filename)


if __name__ == "__main__":
    app.run(debug=True)